#!/usr/bin/env python3
"""
Supplier price-file generator (templator3000).

Reads a supplier price-list workbook and emits the three import files used by
the ERP / POS system, matching the exact column structure, formatting and
calculation logic of the reference TERESA examples:

    PRODUCT_WCEW_<name>.xlsx      product master for the "WCEW" company
    PRODUCT_LIGHTING_<name>.xlsx  product master for the "LIGHTING" company
    SUPPLIER_DETAILS_<name>.xlsx  supplier / product-supplier reconciliation

Calculation logic (verified 1:1 against the reference examples):

    base            = source "Our price (highest of the two)" column
    WCEW cost       = base                (RAW, unrounded  -- see note below)
    LIGHTING cost   = base * 1.10
    LIGHTING trade  = base * 1.25
    LIGHTING retail = base * 1.45
    LIGHTING CPOS   = ceil(base * 1.45 * 1.23)     (round UP to whole number)
    LIGHTING PriceA = base * 1.35
    SUPPLIER price  = base * 1.10

Note on WCEW rounding: the original written spec called for
round_up_half(base) on the WCEW cost, but every reference example row uses the
raw unrounded base (18.95, not 19.0). We follow the examples. Set
WCEW_ROUND_HALF = True to switch to the written-spec behaviour instead.
"""
from __future__ import annotations

import argparse
import math
import os
import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Font


# ----------------------------------------------------------------------------
# Rounding helpers
# ----------------------------------------------------------------------------
def round_up_half(x: float) -> float:
    """Round UP to the nearest 0.5."""
    return math.ceil(x * 2) / 2


def round_up_whole(x: float) -> int:
    """Round UP to the nearest whole number."""
    return int(math.ceil(x))


# ----------------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------------
# If True, WCEW Current Cost Price = round_up_half(base) (written spec).
# If False (default), WCEW cost = base unrounded (matches reference examples).
WCEW_ROUND_HALF = False

TRADE_MULT = 1.25
RETAIL_MULT = 1.45
COST_UPLIFT_MULT = 1.10   # "+10%"
PRICE_A_MULT = 1.35
VAT_MULT = 1.23           # 23% VAT used for CPOS inc-VAT

# Column headers, in order, copied exactly from the reference templates.
WCEW_HEADERS = [
    "Product Code", "Product Description", "Product Group Code",
    "Sales Analysis Code", "Sales VAT Code", "Purchase Analysis Code",
    "Purchase VAT Code", "Alternate Code", "Bin Location", "Unit",
    "Current Cost Price (ex VAT)", "Trade Markup %", "Retail Markup %",
    "CPOS Markup %", "Price A", "Price B", "Price C", "Price D", "Price E",
    "Price F", "Price G", "Price H", "Price I", "Price J", "Price K",
    "Price L", "Price M", "Price N",
]
LIGHTING_HEADERS = [
    "Product Code", "Product Description", "Product Group Code",
    "Sales Analysis Code", "Sales VAT Code", "Purchase Analysis Code",
    "Purchase VAT Code", "Bin Location", "Unit",
    "Current Cost Price (ex VAT)", "Trade Price (ex VAT)",
    "Retail Price (ex VAT)", "CPOS Price (inc VAT)", "Price A", "Price B",
    "Price C", "Price D", "Price E", "Price F", "Price G", "Price H",
    "Price I", "Price J", "Price K", "Price L", "Price M", "Price N",
]
SUPPLIER_HEADERS = [
    "Product Code", "Supplier", "Supplier Product Code", "Price",
    "Last Purchased Date",
]

# Column widths captured from the reference templates (col letter -> width).
WCEW_WIDTHS = {"A": 31.9, "B": 63.0, "C": 19.1, "D": 18.6, "E": 14.7,
               "F": 22.1, "G": 18.3, "H": 19.1, "I": 11.7, "J": 4.9,
               "K": 25.3, "L": 15.4, "M": 15.6, "N": 15.1}
LIGHTING_WIDTHS = {"A": 34.7, "B": 63.0, "C": 19.1, "D": 18.6, "E": 14.7,
                   "F": 22.1, "G": 18.3, "H": 11.7, "I": 23.6, "J": 25.3,
                   "K": 19.1, "L": 19.3}
SUPPLIER_WIDTHS = {"A": 34.7, "B": 14.7, "C": 34.7, "D": 10.3, "E": 18.9}


@dataclass
class SupplierConfig:
    """Per-supplier settings that are not derivable from the price list."""
    brand: str                       # Bin Location, e.g. "FUMAGALLI"
    supplier_code: str               # e.g. "W001"
    product_prefix: str = "W/"       # prefix for SUPPLIER/LIGHTING product code
    # Optional description transform (returns the suffix to append, or "").
    cct_suffix: "callable" = field(default=lambda code, desc: "")


def teresa_cct_suffix(code: str, desc: str) -> str:
    """TERESA-specific: code ending ...Z1L -> ' - 4K', ...Z1R -> ' - 3K'."""
    if re.search(r"Z1L(\b|-|$)", code or ""):
        return " - 4K"
    if re.search(r"Z1R(\b|-|$)", code or ""):
        return " - 3K"
    return ""


# ----------------------------------------------------------------------------
# Source reading
# ----------------------------------------------------------------------------
PRICE_COL_PRIORITY = [
    "our price",          # TERESA: "Our price (highest of the two)"
    "trade price", "rrp", "cost", "price",
]
CODE_ALIASES = ["our product code", "product code", "code", "sku"]
DESC_ALIASES = ["short description", "description", "desc"]
ORDER_ALIASES = ["order code", "supplier code", "alternate code"]


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower()) if s is not None else ""


def find_header_row(ws, max_scan: int = 15):
    """Return (row_index, {normalized_header: col_index})."""
    best = None
    for r in range(1, min(ws.max_row, max_scan) + 1):
        headers = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                headers[_norm(v)] = c
        # A header row should contain a code-ish and a price-ish column.
        has_code = any(a in h for a in CODE_ALIASES for h in headers)
        has_price = any(p in h for p in PRICE_COL_PRIORITY for h in headers)
        if has_code and has_price:
            return r, headers
        if best is None or len(headers) > len(best[1]):
            best = (r, headers)
    return best


def _match_col(headers: dict, aliases: list[str]):
    for a in aliases:
        for h, c in headers.items():
            if a == h:
                return c
    for a in aliases:
        for h, c in headers.items():
            if a in h:
                return c
    return None


def find_price_col(headers: dict):
    for key in PRICE_COL_PRIORITY:
        for h, c in headers.items():
            if key in h:
                return c, h
    return None, None


def read_source(path: str, sheet: str | None = None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    hdr_row, headers = find_header_row(ws)
    code_c = _match_col(headers, CODE_ALIASES)
    desc_c = _match_col(headers, DESC_ALIASES)
    order_c = _match_col(headers, ORDER_ALIASES)
    price_c, price_name = find_price_col(headers)
    if code_c is None or price_c is None:
        raise SystemExit(
            f"Could not identify code/price columns in {path!r}. "
            f"Headers found: {sorted(headers)}"
        )

    rows, skipped = [], []
    for r in range(hdr_row + 1, ws.max_row + 1):
        code = ws.cell(r, code_c).value
        price = ws.cell(r, price_c).value
        if code in (None, ""):
            continue
        if not isinstance(price, (int, float)):
            skipped.append((code, f"non-numeric price {price!r}"))
            continue
        rows.append({
            "code": str(code).strip(),
            "desc": (ws.cell(r, desc_c).value if desc_c else "") or "",
            "order": (ws.cell(r, order_c).value if order_c else "") or "",
            "base": float(price),
        })
    return rows, price_name, ws.title, skipped


# ----------------------------------------------------------------------------
# Writers
# ----------------------------------------------------------------------------
NEG = -100
BOLD = Font(bold=True)


def _write_headers(ws, headers, widths):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(1, c, h)
        cell.font = BOLD
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _txt(ws, r, c, val):
    """Write a value as text ('@' format), preserving leading zeros."""
    cell = ws.cell(r, c, val)
    cell.number_format = "@"
    return cell


def build_wcew(rows, cfg: SupplierConfig):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PRODUCT TEMPLATE"
    _write_headers(ws, WCEW_HEADERS, WCEW_WIDTHS)
    for i, row in enumerate(rows, start=2):
        base = row["base"]
        cost = round_up_half(base) if WCEW_ROUND_HALF else base
        desc = str(row["desc"]) + cfg.cct_suffix(row["code"], row["desc"])
        _txt(ws, i, 1, row["code"])
        ws.cell(i, 2, desc)
        _txt(ws, i, 3, "077")
        _txt(ws, i, 4, "001")
        _txt(ws, i, 5, "3")
        _txt(ws, i, 6, "001")
        _txt(ws, i, 7, "3")
        _txt(ws, i, 8, str(row["order"]))
        _txt(ws, i, 9, cfg.brand)
        _txt(ws, i, 10, "EA")
        ws.cell(i, 11, cost)                       # Current Cost Price
        ws.cell(i, 12, 25)                         # Trade Markup %
        ws.cell(i, 13, 45)                         # Retail Markup %
        ws.cell(i, 14, 45)                         # CPOS Markup %
        ws.cell(i, 15, 35)                         # Price A (markup %)
        ws.cell(i, 16, 20)                         # Price B (markup %)
        for c in range(17, 22):                    # Price C..G
            ws.cell(i, c, NEG)
        ws.cell(i, 22, 10)                         # Price H (markup %)
        for c in range(23, 29):                    # Price I..N
            ws.cell(i, c, NEG)
    return wb


def build_lighting(rows, cfg: SupplierConfig):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PRODUCT TEMPLATE"
    _write_headers(ws, LIGHTING_HEADERS, LIGHTING_WIDTHS)
    ws.freeze_panes = "B1"
    for i, row in enumerate(rows, start=2):
        base = row["base"]
        desc = str(row["desc"]) + cfg.cct_suffix(row["code"], row["desc"])
        _txt(ws, i, 1, cfg.product_prefix + row["code"])
        ws.cell(i, 2, desc)
        _txt(ws, i, 3, "009")
        _txt(ws, i, 4, "001")
        _txt(ws, i, 5, "3")
        _txt(ws, i, 6, "001")
        _txt(ws, i, 7, "3")
        _txt(ws, i, 8, cfg.brand)
        _txt(ws, i, 9, "EA")
        ws.cell(i, 10, base * COST_UPLIFT_MULT).number_format = "0.00"
        ws.cell(i, 11, base * TRADE_MULT).number_format = "0.00"
        ws.cell(i, 12, base * RETAIL_MULT).number_format = "0.00"
        ws.cell(i, 13, round_up_whole(base * RETAIL_MULT * VAT_MULT))  # CPOS
        ws.cell(i, 14, base * PRICE_A_MULT).number_format = "0.0"      # Price A
        for c in range(15, 28):                    # Price B..N
            ws.cell(i, c, NEG)
    return wb


def build_supplier(rows, cfg: SupplierConfig):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ProdSuppRecImport"
    _write_headers(ws, SUPPLIER_HEADERS, SUPPLIER_WIDTHS)
    for i, row in enumerate(rows, start=2):
        base = row["base"]
        ws.cell(i, 1, cfg.product_prefix + row["code"])
        _txt(ws, i, 2, cfg.supplier_code)
        ws.cell(i, 3, row["code"])
        ws.cell(i, 4, base * COST_UPLIFT_MULT).number_format = "#,##0.00"
        d = ws.cell(i, 5, "=TODAY()")
        d.number_format = "mm-dd-yy"
    return wb


# ----------------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------------
def generate(source, outdir, cfg: SupplierConfig, name=None, sheet=None):
    rows, price_name, sheet_title, skipped = read_source(source, sheet)
    stem = name or os.path.splitext(os.path.basename(source))[0]
    os.makedirs(outdir, exist_ok=True)
    outputs = {
        f"PRODUCT_WCEW_{stem}.xlsx": build_wcew(rows, cfg),
        f"PRODUCT_LIGHTING_{stem}.xlsx": build_lighting(rows, cfg),
        f"SUPPLIER_DETAILS_{stem}.xlsx": build_supplier(rows, cfg),
    }
    paths = []
    for fname, wb in outputs.items():
        p = os.path.join(outdir, fname)
        wb.save(p)
        paths.append(p)
    return paths, rows, price_name, sheet_title, skipped


def main():
    ap = argparse.ArgumentParser(description="Generate ERP/POS price files.")
    ap.add_argument("source", help="source price-list .xlsx")
    ap.add_argument("-o", "--outdir", default=".", help="output directory")
    ap.add_argument("-n", "--name", help="stem for output filenames")
    ap.add_argument("--sheet", help="source sheet name (default: first)")
    ap.add_argument("--brand", required=True, help="Bin Location / brand name")
    ap.add_argument("--supplier-code", required=True, help="e.g. W001")
    ap.add_argument("--prefix", default="W/", help="product-code prefix")
    ap.add_argument("--teresa-cct", action="store_true",
                    help="apply TERESA CCT suffix (Z1L->4K, Z1R->3K)")
    args = ap.parse_args()

    cfg = SupplierConfig(
        brand=args.brand,
        supplier_code=args.supplier_code,
        product_prefix=args.prefix,
        cct_suffix=teresa_cct_suffix if args.teresa_cct else (lambda c, d: ""),
    )
    paths, rows, price_name, sheet_title, skipped = generate(
        args.source, args.outdir, cfg, name=args.name, sheet=args.sheet)

    print(f"Source sheet:     {sheet_title}")
    print(f"Price column:     {price_name!r}")
    print(f"Rows processed:   {len(rows)}")
    print(f"Rows skipped:     {len(skipped)}")
    for code, why in skipped:
        print(f"   - {code}: {why}")
    print("Files written:")
    for p in paths:
        print(f"   {p}")


if __name__ == "__main__":
    main()
